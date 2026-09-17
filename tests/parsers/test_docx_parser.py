"""Tests for the DOCX document parser."""

from pathlib import Path

from docx import Document

from resumeoptimizer.parsers import DOCXParser


def test_docx_parser_reads_document(tmp_path: Path) -> None:
    """Test that DOCXParser extracts paragraph text."""
    file_path = tmp_path / "resume.docx"

    document = Document()
    document.add_paragraph("Python Developer")
    document.add_paragraph("Machine Learning Engineer")
    document.save(file_path)

    parser = DOCXParser()

    result = parser.parse(file_path)

    assert "Python Developer" in result
    assert "Machine Learning Engineer" in result
